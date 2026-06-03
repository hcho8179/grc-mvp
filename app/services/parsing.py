"""File parsing service: PDF, DOCX, XLSX."""

import io
from pathlib import Path


def parse_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except PdfReadError as exc:
        raise ValueError(f"Could not read PDF: {exc}") from exc
    except Exception as exc:
        raise ValueError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        raise ValueError("PDF is password-protected and cannot be read.")

    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)

    return "\n\n".join(parts)


def parse_docx(file_bytes: bytes) -> str:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Could not read DOCX: {exc}") from exc

    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _xlsx_hidden_rows(ws) -> set[int]:
    return {r for r, rd in ws.row_dimensions.items() if rd.hidden}


def _preview_worksheet(ws) -> dict:
    hidden_rows = _xlsx_hidden_rows(ws)
    rows = list(ws.iter_rows(values_only=True))
    sheet_name = ws.title

    if not rows:
        return {"columns": [], "sample_rows": [], "total_rows": 0, "sheet_name": sheet_name}

    columns = [
        str(cell) if cell is not None else f"Column {i + 1}"
        for i, cell in enumerate(rows[0])
    ]
    data_rows = [
        row for i, row in enumerate(rows[1:], start=2)
        if i not in hidden_rows
    ]
    sample = [
        [str(cell) if cell is not None else "" for cell in row]
        for row in data_rows[:5]
    ]
    return {
        "columns": columns,
        "sample_rows": sample,
        "total_rows": len(data_rows),
        "sheet_name": sheet_name,
    }


def parse_xlsx_questionnaire_sheets_preview(file_bytes: bytes) -> dict:
    """Return a preview for every worksheet (multi-tab workbooks)."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read XLSX: {exc}") from exc

    sheets: list[dict] = []
    for index, name in enumerate(wb.sheetnames):
        preview = _preview_worksheet(wb[name])
        preview["index"] = index
        preview["name"] = name
        sheets.append(preview)
    wb.close()

    if not sheets:
        raise ValueError("Workbook contains no worksheets.")

    return {"sheets": sheets, "default_sheet_index": 0}


def parse_xlsx_questionnaire_preview(file_bytes: bytes, sheet_index: int = 0) -> dict:
    """Return column headers, sample rows, total row count, and sheet name for one sheet."""
    data = parse_xlsx_questionnaire_sheets_preview(file_bytes)
    sheets = data["sheets"]
    if sheet_index < 0 or sheet_index >= len(sheets):
        raise ValueError(f"Invalid sheet index {sheet_index}.")
    return sheets[sheet_index]


def parse_xlsx_qa(file_bytes: bytes, question_col: int, answer_col: int, sheet_index: int = 0) -> list[dict]:
    """Parse an XLSX file into Q&A pairs using the given 0-based column indices."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read XLSX: {exc}") from exc

    if sheet_index < 0 or sheet_index >= len(wb.sheetnames):
        wb.close()
        raise ValueError(f"Invalid sheet index {sheet_index}.")
    ws = wb[wb.sheetnames[sheet_index]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    required_cols = max(question_col, answer_col)
    pairs: list[dict] = []
    for row in rows[1:]:
        if len(row) <= required_cols:
            continue
        question = str(row[question_col]).strip() if row[question_col] is not None else ""
        answer = str(row[answer_col]).strip() if row[answer_col] is not None else ""
        if question:
            pairs.append({"question": question, "answer": answer})

    return pairs


def parse_xlsx_questionnaire_full(
    file_bytes: bytes,
    question_col: int,
    answer_col: int | None = None,
    sheet_index: int = 0,
) -> list[dict]:
    """Parse a blank questionnaire XLSX into structured question rows."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read XLSX: {exc}") from exc

    if sheet_index < 0 or sheet_index >= len(wb.sheetnames):
        wb.close()
        raise ValueError(f"Invalid sheet index {sheet_index}.")

    ws = wb[wb.sheetnames[sheet_index]]
    hidden_rows = _xlsx_hidden_rows(ws)
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    results: list[dict] = []
    for idx, row in enumerate(rows[1:]):
        xlsx_row = idx + 2
        if xlsx_row in hidden_rows:
            continue
        if len(row) <= question_col:
            continue
        question = str(row[question_col]).strip() if row[question_col] is not None else ""
        if not question:
            continue
        existing_answer: str | None = None
        if answer_col is not None and len(row) > answer_col:
            val = row[answer_col]
            existing_answer = str(val).strip() if val is not None else None
            if not existing_answer:
                existing_answer = None
        results.append({"row_index": idx, "question": question, "existing_answer": existing_answer})

    return results


def parse_xlsx_preview(file_bytes: bytes, sheet_index: int = 0) -> dict:
    """Return column headers, sample rows, and total row count for the column-picker UI."""
    return parse_xlsx_questionnaire_preview(file_bytes, sheet_index=sheet_index)


def parse_file(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the appropriate text parser based on file extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_bytes)
    if ext == ".docx":
        return parse_docx(file_bytes)
    if ext == ".xlsx":
        raise ValueError("XLSX files require column selection — use parse_xlsx_* functions directly.")
    raise ValueError(f"Unsupported file type: {ext!r}. Supported: .pdf, .docx, .xlsx")
