#!/usr/bin/env python
"""
Add additional knowledge base documents without wiping existing data.

Usage (from project root):
    python scripts/add_more_docs.py
"""

import asyncio
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document as DocxDocument

from app.db import SessionLocal
from app.models import Document, FileType, KnowledgeChunk, SourceType
from app.services.chunking import chunk_text
from app.services.embeddings import get_embedding_provider
from app.services.parsing import parse_docx

# ---------------------------------------------------------------------------
# New policy documents
# ---------------------------------------------------------------------------

CLOUD_SECURITY_POLICY = """CLOUD SECURITY POLICY
Version 1.4 | Effective: March 1, 2026 | Owner: Platform Security Team

1. PURPOSE AND SCOPE
This policy governs Acme Corp's use of cloud infrastructure and services. It applies to all AWS services, SaaS tools, and cloud-based development environments used by employees and contractors.

2. INFRASTRUCTURE AS CODE
All cloud infrastructure is defined and provisioned using Terraform. Manual console changes to production environments are prohibited except during declared incidents. All Terraform changes go through pull request review and apply only via CI/CD pipelines using short-lived OIDC credentials.

3. IDENTITY AND ACCESS MANAGEMENT
AWS IAM roles follow least-privilege principles. No IAM users with long-lived access keys are created. All AWS access uses role assumption with session durations capped at 1 hour. AWS Organizations SCPs enforce guardrails across all accounts, preventing disabling of CloudTrail, GuardDuty, and Security Hub.

Cross-account access uses IAM role chaining with explicit trust policies. External access to AWS is via AWS SSO integrated with Okta for centralized identity management. AWS root account credentials are stored in a hardware security module and used only for break-glass scenarios.

4. NETWORK ARCHITECTURE
Production workloads run in private subnets with no direct internet ingress. Internet-facing services use Application Load Balancers in public subnets. All inter-service communication uses AWS PrivateLink or VPC peering — no public internet transit for internal traffic.

Network segmentation is enforced via Security Groups and NACLs. A default-deny posture applies to all new services. Security Group rules are reviewed quarterly via automated policy-as-code checks.

5. SECRETS MANAGEMENT
All secrets, API keys, and credentials are stored in AWS Secrets Manager or AWS Parameter Store (SecureString). Hardcoded secrets in source code are detected automatically via Trufflehog and git-secrets pre-commit hooks. Detected violations block merge and trigger immediate rotation.

Database credentials rotate automatically every 30 days via Secrets Manager rotation. Application code retrieves secrets at runtime — secrets are never baked into container images or environment variables in task definitions.

6. CONTAINER AND IMAGE SECURITY
All container images are built from approved base images maintained by the Platform team. Images are scanned for vulnerabilities via Amazon ECR image scanning and Trivy in CI. Images with critical CVEs are blocked from deployment. Base images are rebuilt weekly with the latest OS patches.

Containers run as non-root users. Privileged container mode is disabled. Read-only root filesystems are enforced where possible. Resource limits (CPU and memory) are defined for all containers.

7. DATA ENCRYPTION IN THE CLOUD
All S3 buckets enforce server-side encryption with AWS KMS (SSE-KMS). S3 bucket public access is blocked at the account level. CloudFront distributions use HTTPS-only with TLS 1.2 minimum. RDS instances use encryption at rest with KMS keys and encrypted storage volumes.

DynamoDB tables are encrypted at rest. ElastiCache clusters use in-transit and at-rest encryption. All EBS volumes attached to EC2 instances are encrypted. Unencrypted resource creation is blocked via AWS Config rules and SCPs.

8. LOGGING AND OBSERVABILITY
CloudTrail is enabled in all regions with log file validation and delivered to a dedicated, tamper-resistant S3 bucket with Object Lock enabled. AWS Config records all resource configuration changes. VPC Flow Logs are enabled for all production VPCs.

GuardDuty is enabled in all regions and accounts. Security Hub aggregates findings from GuardDuty, Inspector, and Macie. All findings above Medium severity trigger PagerDuty alerts to the Security Operations team within 5 minutes.

9. DISASTER RECOVERY
Primary region: AWS us-east-1. DR region: AWS us-west-2. RDS uses Multi-AZ with automated failover and cross-region read replicas promoted during DR events. S3 data is replicated cross-region via S3 Cross-Region Replication with versioning enabled. Recovery procedures are documented and tested annually.
"""

NETWORK_SECURITY_POLICY = """NETWORK SECURITY POLICY
Version 2.0 | Effective: February 1, 2026 | Owner: Network Security Team

1. PURPOSE
This policy defines controls for Acme Corp's network infrastructure to protect against unauthorized access, data interception, and network-based attacks.

2. NETWORK SEGMENTATION
The network is divided into security zones: Production, Staging, Development, Corporate, and DMZ. Traffic between zones is controlled by firewalls with explicit allow rules. Default-deny rules apply at all zone boundaries. Production environments are isolated from development and corporate networks.

Database servers reside in dedicated database subnets with no direct access from the internet. Application servers communicate with databases only through defined application-layer connections on specific ports. Direct database connections from developer workstations require VPN plus jump host access, which is logged and time-limited.

3. FIREWALL MANAGEMENT
Firewalls are configured using infrastructure-as-code. All rule changes require change management approval and peer review. Unused firewall rules are removed quarterly. Firewall rule sets are reviewed annually by an independent security team. AWS Security Groups serve as host-based firewalls with stateful inspection.

4. VPN AND REMOTE ACCESS
Corporate VPN uses WireGuard with certificate-based authentication integrated with Okta. Split-tunnel VPN is not used for production access — all production traffic routes through the VPN. VPN sessions require MFA and have a maximum idle timeout of 30 minutes. VPN access logs are retained for 12 months.

5. INTRUSION DETECTION AND PREVENTION
AWS GuardDuty provides network-level threat detection for all VPCs. Network traffic baselines are established and anomalies trigger automated alerts. Known malicious IP ranges are blocked at the network perimeter and updated daily via threat intelligence feeds integrated with AWS Network Firewall.

6. DNS SECURITY
DNS queries from corporate workstations use encrypted DNS-over-HTTPS (DoH) resolvers. Internal DNS is separated from external DNS. DNSSEC is enabled for the acme.com domain. DNS query logs are collected and analyzed for indicators of compromise, including DNS tunneling patterns.

7. DDoS PROTECTION
AWS Shield Advanced is deployed across all public-facing endpoints providing volumetric and protocol attack mitigation. AWS WAF protects application layer with rules updated automatically via managed rule groups. Rate limiting is applied at the API Gateway and CloudFront layers. DDoS response runbooks are tested in annual tabletop exercises.

Layer 7 protection rules include: rate limiting per IP (1000 requests/minute), geo-blocking for regions without customers, bot detection, and SQL injection/XSS prevention. WAF logs are retained for 90 days and analyzed in the SIEM.

8. WIRELESS SECURITY
Corporate wireless networks use WPA3-Enterprise with 802.1X authentication via Okta RADIUS. Guest networks are isolated with no access to corporate resources. Rogue access point detection is deployed in all offices. Wireless network traffic is encrypted end-to-end — the wireless network is treated as untrusted infrastructure.

9. NETWORK MONITORING
NetFlow/VPC Flow Logs are collected for all network traffic. Logs are analyzed in real-time by the SIEM for anomalies, lateral movement, and data exfiltration patterns. Bandwidth anomalies exceeding 3 standard deviations from baseline trigger automated investigation tickets. Network monitoring data is retained for 12 months.
"""

SECURE_DEVELOPMENT_POLICY = """SECURE SOFTWARE DEVELOPMENT LIFECYCLE (SSDLC) POLICY
Version 1.2 | Effective: January 15, 2026 | Owner: Engineering Security

1. PURPOSE
This policy defines security requirements integrated into Acme Corp's software development lifecycle to ensure secure, reliable software from design through deployment and decommission.

2. SECURITY REQUIREMENTS
All new features undergo a security requirements review during the design phase. Threat modeling using STRIDE methodology is required for features that process personal data, handle authentication, or expose new API endpoints. Threat models are documented in Confluence and linked to the associated Jira epic.

High-risk features (payment processing, authentication changes, data export) require a dedicated Application Security review before development begins. Security requirements are tracked as acceptance criteria in Jira tickets and must be verified before the ticket is closed.

3. SECURE CODING STANDARDS
Developers follow OWASP Secure Coding Practices. Code reviews are mandatory for all changes — no self-merge is permitted. Reviews must include at least one senior engineer who evaluates security implications. Security-sensitive code (authentication, cryptography, authorization) requires review by a member of the Application Security team.

Input validation is required at all trust boundaries. Parameterized queries are used for all database operations — string concatenation in SQL queries is a blocking code review finding. Output encoding is applied for all user-controlled data rendered in HTML, JSON, or other formats.

4. STATIC APPLICATION SECURITY TESTING (SAST)
Semgrep SAST scanning runs on every pull request with Acme Corp's custom rule set. Critical and high findings block merge. Medium findings must be acknowledged with a security comment before merge. SAST results are tracked in the Application Security dashboard with SLA requirements: critical within 72 hours, high within 7 days.

CodeQL analysis runs on all repositories weekly with results reviewed by Application Security. Third-party code introduced via dependencies is scanned using OWASP Dependency-Check and Snyk. License compliance is verified automatically for all new dependencies.

5. DYNAMIC APPLICATION SECURITY TESTING (DAST)
OWASP ZAP automated DAST scanning runs against the staging environment on every deploy. The authenticated scan tests all API endpoints using service account credentials. DAST findings above Medium severity must be remediated before production deployment. Results are reviewed weekly by Application Security.

6. SECRETS DETECTION
Trufflehog and git-secrets pre-commit hooks prevent accidental secret commits. GitHub Advanced Security secret scanning monitors all repositories continuously. Any detected secret is treated as compromised and rotated within 1 hour regardless of where it was found. Commit history is audited and cleaned if secrets are discovered.

7. DEPENDENCY MANAGEMENT
All dependencies are pinned to specific versions using lock files (Poetry lock for Python, package-lock.json for Node). Dependabot monitors all repositories and creates PRs for security updates within 24 hours of CVE disclosure. Critical dependency CVEs are treated as P1 issues with 72-hour remediation SLA.

Software Bill of Materials (SBOM) is generated for every production release using syft. SBOMs are stored and available to enterprise customers upon request. Container base images are updated weekly in an automated pipeline.

8. SECURITY TESTING IN CI/CD
The CI pipeline includes: SAST scan, dependency vulnerability check, container image scan, secrets detection, and unit tests with security test cases. All checks must pass for deployment to staging. Production deployments additionally require DAST scan completion and sign-off from an on-call engineer.

Deployment to production uses GitOps via ArgoCD with approval gates. Deployments are audited with full traceability from commit to production. Rollback capability is maintained for all services with automated rollback triggered by error rate SLOs.

9. POST-PRODUCTION SECURITY
Production systems are monitored by AWS Inspector for OS and application vulnerabilities. New critical CVEs trigger automated ticket creation and notification to the responsible team. Runtime application self-protection (RASP) monitors for anomalous behavior in production API handlers.
"""

EXTRA_QA_PAIRS_2 = [
    ("Do you support single sign-on (SSO)?",
     "Yes, Acme Corp supports SAML 2.0 and OIDC-based SSO. We integrate with major identity providers including Okta, Azure AD, and Google Workspace. SSO can be enforced at the organization level so that password-based login is disabled for all users. Just-in-time provisioning is supported via SCIM 2.0."),

    ("How do you handle employee offboarding from a security perspective?",
     "Upon termination, all system access is revoked within 1 hour for voluntary departures and immediately upon HR notification for involuntary terminations. This includes SSO, VPN, cloud consoles, GitHub, and SaaS tools. Company devices are remotely wiped and recovered. Service account credentials associated with the employee are rotated within 24 hours. Offboarding checklists are tracked in our ITSM system."),

    ("What certifications does Acme Corp hold?",
     "Acme Corp holds SOC 2 Type II (annual audit, most recent period Oct 2024–Sep 2025), ISO 27001 certification (certified November 2024, valid through 2027), and is GDPR and CCPA compliant. We are currently pursuing FedRAMP Moderate authorization with an anticipated authorization date of Q3 2026. Reports are available under NDA to enterprise customers upon request."),

    ("Do you support role-based access control (RBAC) within your platform?",
     "Yes, the platform supports granular RBAC with predefined roles (Admin, Editor, Viewer, Auditor) and custom role creation. Permissions are enforced at the API layer on every request. Administrators can assign roles at organization, team, or individual resource levels. All permission changes are logged in the audit trail and can be exported for compliance reviews."),

    ("How do you ensure high availability and uptime?",
     "Our infrastructure is deployed across multiple AWS Availability Zones with automatic failover. We maintain a 99.9% uptime SLA for our production API and dashboard. Redundant load balancers, auto-scaling groups, and multi-AZ RDS instances eliminate single points of failure. Our historical uptime over the past 12 months is 99.96%, tracked on our public status page at status.acme.com."),

    ("What is your process for customer data deletion upon contract termination?",
     "Upon contract termination, customer data is soft-deleted immediately (no longer accessible via the product) and permanently deleted from all storage within 30 days, including backups that expire on their normal schedule. We provide a written confirmation of deletion within 45 days. Customers may request an accelerated deletion within 7 business days. We do not retain any customer data post-deletion except as required by law."),

    ("Do you have a formal risk assessment process?",
     "Yes, Acme Corp conducts formal risk assessments annually using the NIST Cybersecurity Framework. The risk register is reviewed and updated quarterly by the Security team. Risk owners are assigned for each identified risk. The CISO presents the risk register and treatment status to the Board of Directors quarterly. Material risks are disclosed in customer-facing security documentation."),

    ("How do you handle zero-day vulnerabilities?",
     "Upon disclosure of a zero-day affecting our stack, the Security team convenes within 2 hours to assess exposure. If affected, a P1 incident is declared and mitigation (WAF rule, network isolation, or patch) is applied within 24 hours. Customers are notified proactively if the vulnerability affects systems that process their data. We subscribe to vendor security advisories and threat intelligence feeds to minimize detection lag."),

    ("What is your patch management process?",
     "OS patches are applied to all production instances within 7 days of release for security patches and 30 days for non-security patches. Critical patches (CVSS 9.0+) are applied within 72 hours. Patching is automated via AWS Systems Manager Patch Manager with a defined patch baseline. Container base images are rebuilt weekly with the latest OS packages. Patch compliance is reported weekly to the Security team."),

    ("Do you have a formal asset inventory?",
     "Yes, we maintain a comprehensive asset inventory covering all cloud resources (via AWS Config), endpoints (via MDM), software (via SBOM), and SaaS tools (via Okta integrations). The inventory is updated in real-time for cloud resources and weekly for endpoints. Asset owners are assigned and responsible for ensuring assets meet security baselines. Unmanaged assets detected by network scanning trigger automated alerts."),

    ("How do you manage API security?",
     "All APIs require authentication via OAuth 2.0 Bearer tokens or API keys with scoped permissions. API keys are rotated annually and can be revoked instantly. Rate limiting is enforced at the API Gateway (1000 requests/minute per key by default). All API traffic is logged with request/response metadata (excluding bodies containing PII). OWASP API Security Top 10 is part of our API design review checklist and DAST test suite."),

    ("Do you perform code reviews for security?",
     "Yes, all code changes require peer review before merging — self-merge is prohibited. Security-sensitive changes (authentication, authorization, cryptography, data handling) require review by the Application Security team in addition to standard peer review. SAST results are visible in the pull request and critical findings block merge. We conduct targeted security-focused reviews for all major features before release."),

    ("What is your approach to data minimization?",
     "Acme Corp applies data minimization principles throughout product design. New features undergo a Privacy Impact Assessment to identify and eliminate unnecessary data collection. We do not collect behavioral analytics beyond what is needed for product functionality and security monitoring. Data fields are reviewed annually for continued necessity. Customers can configure data retention periods shorter than our maximum within the platform settings."),

    ("How do you handle security for mobile applications?",
     "Our mobile applications (iOS and Android) are developed following OWASP Mobile Security Testing Guide. Certificate pinning prevents MITM attacks. Sensitive data is stored only in platform-encrypted keystores — never in plain text on device. The apps undergo annual mobile penetration testing by a third party. Application binary protection (obfuscation and anti-tampering) is applied to release builds."),

    ("Do you have a formal change management process for security-impacting changes?",
     "Yes, all changes to production systems go through our change management process: change request creation, risk assessment, peer review, approval (standard or emergency), and post-change validation. Emergency changes require post-facto review within 24 hours. Security-impacting changes (firewall rules, IAM policies, cryptographic settings) require Security team approval before implementation. All changes are logged with full audit trail."),

    ("How do you manage sub-processor security?",
     "All sub-processors handling customer personal data are listed at acme.com/sub-processors and are reviewed annually for SOC 2 or equivalent certification. Before onboarding a new sub-processor, a vendor security assessment is completed including review of their security policies, certifications, and breach history. Customers are notified 30 days before adding new sub-processors and may object. Sub-processors sign Data Processing Agreements with security requirements equivalent to our own obligations."),

    ("What is your SLA for responding to security vulnerabilities reported by customers?",
     "We acknowledge all customer-reported security vulnerabilities within 24 hours. Critical vulnerabilities are triaged within 4 hours and remediated within 72 hours with proactive customer notification. High vulnerabilities are remediated within 7 days. We operate a private bug bounty program on HackerOne where researchers can report vulnerabilities. We do not pursue legal action against good-faith security researchers."),

    ("How do you ensure the security of your CI/CD pipeline?",
     "CI/CD pipelines run in isolated environments with no persistent credentials. Short-lived OIDC tokens authenticate to AWS and other services. Pipeline configurations are stored in version control and require pull request review. Production deployments require additional approval from an on-call engineer and cannot be triggered by automated processes alone. Pipeline logs are retained for 90 days and monitored for anomalies."),

    ("Do you encrypt backups?",
     "Yes, all backups are encrypted using AES-256 with keys managed in AWS KMS. RDS automated backups are encrypted by default using the same KMS key as the source database. S3 backups use SSE-KMS. Backup encryption keys are rotated annually. Backup restoration is tested quarterly, including validation that encrypted backups can be successfully decrypted and restored within our RTO of 4 hours."),

    ("What access does Acme Corp staff have to customer data?",
     "Acme Corp staff do not have standing access to customer data. Access to production data requires a formal request, manager approval, and Security team sign-off. All access is time-limited (maximum 8 hours), logged with justification, and audited monthly. Customer support staff access only the minimum data necessary to resolve the support ticket. Access to production databases requires jump host with session recording."),
]

EXTRA_QA_PAIRS = [
    ("How do you manage cloud infrastructure access and prevent privilege escalation?",
     "All AWS access uses IAM roles with OIDC-based short-lived credentials (1-hour session max) via AWS SSO integrated with Okta. No long-lived IAM user access keys exist. AWS Organizations SCPs enforce guardrails preventing disabling of CloudTrail, GuardDuty, and Security Hub. Privilege escalation paths are reviewed quarterly via automated IAM policy analysis."),

    ("Describe your container security controls.",
     "All container images are built from Platform-team-maintained base images, scanned via ECR and Trivy in CI, and blocked from deployment if critical CVEs are detected. Containers run as non-root with read-only root filesystems where possible. Privileged mode is disabled. Base images are rebuilt weekly with the latest OS patches."),

    ("How do you protect against SQL injection and other web application attacks?",
     "Parameterized queries are enforced for all database operations — string concatenation in SQL is a blocking code review finding. SAST (Semgrep) scans every pull request for injection vulnerabilities. AWS WAF with managed rule groups protects all public endpoints against OWASP Top 10 attacks. DAST scanning (OWASP ZAP) runs against staging on every deploy."),

    ("Do you use Infrastructure as Code? How are changes controlled?",
     "All infrastructure is defined in Terraform. Manual console changes to production are prohibited except during declared incidents. Terraform applies only via CI/CD pipelines using short-lived OIDC credentials. All changes require pull request review and approval before apply. Drift detection runs daily and alerts on any unmanaged configuration changes."),

    ("What is your process for managing secrets and API keys?",
     "All secrets are stored in AWS Secrets Manager or Parameter Store (SecureString). Hardcoded secrets are detected via Trufflehog and git-secrets pre-commit hooks — violations block merge and trigger immediate rotation. Database credentials rotate automatically every 30 days. Secrets are never baked into container images or task definition environment variables."),

    ("How do you handle network segmentation between production and development environments?",
     "Production, staging, development, and corporate networks are in separate VPCs with no direct peering. Traffic between zones passes through firewalls with explicit allow rules and default-deny posture. Database subnets have no internet access. Developer access to production databases requires VPN plus a time-limited jump host session, which is fully logged."),

    ("What threat modeling process do you follow?",
     "We use STRIDE methodology for all features processing personal data, handling authentication, or exposing new API endpoints. Threat models are documented in Confluence and linked to the Jira epic. High-risk features require Application Security team review before development begins. Threat model findings are tracked as acceptance criteria and verified before closure."),

    ("Do you have a software bill of materials (SBOM)?",
     "Yes, SBOMs are generated for every production release using syft and stored with the release artifacts. SBOMs are available to enterprise customers upon request. Dependabot monitors all repositories for CVEs daily. Critical dependency vulnerabilities have a 72-hour remediation SLA. A complete dependency inventory is maintained for all services."),

    ("How do you protect DNS and prevent DNS-based attacks?",
     "Corporate DNS uses encrypted DNS-over-HTTPS (DoH). DNSSEC is enabled for the acme.com domain. Internal DNS is separated from external. DNS query logs are analyzed for tunneling and C2 patterns in the SIEM. Recursive DNS resolvers are restricted to known corporate IP ranges. DNS configuration changes require change management approval."),

    ("What runtime monitoring do you have in production?",
     "AWS GuardDuty monitors all VPCs and accounts for threat patterns. AWS Inspector continuously scans production systems for OS and application vulnerabilities. New critical CVEs trigger automated ticket creation. CloudTrail with Object Lock protects all API audit logs from tampering. Security Hub aggregates findings and alerts the Security Operations team on anything above Medium severity within 5 minutes."),
]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

async def main() -> None:
    db = SessionLocal()
    provider = get_embedding_provider()

    new_policies = [
        ("Acme Corp — Cloud Security Policy.docx", CLOUD_SECURITY_POLICY),
        ("Acme Corp — Network Security Policy.docx", NETWORK_SECURITY_POLICY),
        ("Acme Corp — Secure Development Lifecycle Policy.docx", SECURE_DEVELOPMENT_POLICY),
    ]

    print("📄  Adding new policy documents…")
    for filename, content in new_policies:
        doc_obj = DocxDocument()
        for para in content.strip().split("\n\n"):
            doc_obj.add_paragraph(para.strip())
        buf = io.BytesIO()
        doc_obj.save(buf)
        docx_bytes = buf.getvalue()

        text_content = parse_docx(docx_bytes)
        chunks = chunk_text(text_content)

        print(f"   {filename}: {len(chunks)} chunks — embedding…")

        document = Document(
            filename=filename,
            file_type=FileType.docx,
            chunk_count=len(chunks),
        )
        db.add(document)
        db.flush()

        embeddings = await provider.embed(chunks)
        for chunk_text_val, emb in zip(chunks, embeddings):
            db.add(KnowledgeChunk(
                document_id=document.id,
                source_document=filename,
                source_type=SourceType.policy,
                text=chunk_text_val,
                embedding=emb,
            ))
        db.commit()
        print(f"   ✓ {filename} added")

    all_batches = [
        ("Acme Corp — Extended Q&A Knowledge Base.xlsx", EXTRA_QA_PAIRS),
        ("Acme Corp — Extended Q&A Knowledge Base Vol 2.xlsx", EXTRA_QA_PAIRS_2),
    ]

    print("\n💬  Adding extra Q&A pairs…")
    for filename, pairs in all_batches:
        from app.models import Document as Doc
        qa_doc = Doc(
            filename=filename,
            file_type=FileType.xlsx,
            chunk_count=len(pairs),
        )
        db.add(qa_doc)
        db.flush()

        qa_texts = [f"Q: {q}\nA: {a}" for q, a in pairs]
        embeddings = await provider.embed(qa_texts)
        for (question, answer), emb in zip(pairs, embeddings):
            db.add(KnowledgeChunk(
                document_id=qa_doc.id,
                source_document=qa_doc.filename,
                source_type=SourceType.past_answer,
                text=f"Q: {question}\nA: {answer}",
                embedding=emb,
            ))
        db.commit()
        print(f"   ✓ {len(pairs)} pairs added from {filename}")

    from app.models import KnowledgeChunk as KC
    total = db.query(KC).count()
    print(f"\n✅  Done — knowledge base now has {total} total chunks")
    print("   Regen any question to see updated answers from the new documents.")
    db.close()


if __name__ == "__main__":
    asyncio.run(main())
