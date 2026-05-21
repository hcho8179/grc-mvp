"""Tests for Phase 8 export and template services."""

import io

import openpyxl
import pytest

from app.services.export import export_questionnaire_to_xlsx
from app.services.templates import (
    extract_format_config,
    is_cell_in_merge_interior,
    save_template,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_simple_xlsx(questions: list[str], header: str = "Question") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=header)
    ws.cell(row=1, column=2, value="Answer")
    for i, q in enumerate(questions, start=2):
        ws.cell(row=i, column=1, value=q)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_merged_xlsx() -> bytes:
    """Worksheet with merged header row and questions/answers below."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:B1")
    ws.cell(row=1, column=1, value="Questionnaire")
    ws.cell(row=2, column=1, value="Question")
    ws.cell(row=2, column=2, value="Answer")
    ws.cell(row=3, column=1, value="Do you encrypt data?")
    ws.cell(row=4, column=1, value="Do you have MFA?")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_formula_xlsx() -> bytes:
    """Worksheet where the answer column contains a formula."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Question")
    ws.cell(row=1, column=2, value="Answer")
    ws.cell(row=2, column=1, value="Do you encrypt data?")
    ws.cell(row=2, column=2, value="=A2")  # formula in answer column
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# extract_format_config
# ---------------------------------------------------------------------------

def test_extract_format_config_basic():
    xlsx = _make_simple_xlsx(["Q1", "Q2"])
    config = extract_format_config(xlsx)
    assert "merged_cells" in config
    assert "column_widths" in config
    assert "row_heights" in config
    assert "formula_cells" in config
    assert config["merged_cells"] == []


def test_extract_format_config_merged():
    xlsx = _make_merged_xlsx()
    config = extract_format_config(xlsx)
    assert len(config["merged_cells"]) == 1
    assert "A1:B1" in config["merged_cells"]


def test_extract_format_config_formulas():
    xlsx = _make_formula_xlsx()
    config = extract_format_config(xlsx)
    assert "B2" in config["formula_cells"]


# ---------------------------------------------------------------------------
# is_cell_in_merge_interior
# ---------------------------------------------------------------------------

def test_merge_interior_detection():
    format_config = {"merged_cells": ["A1:C1"]}
    # A1 is top-left → not interior
    assert not is_cell_in_merge_interior(1, 1, format_config)
    # B1 and C1 are interior
    assert is_cell_in_merge_interior(1, 2, format_config)
    assert is_cell_in_merge_interior(1, 3, format_config)
    # Row 2 is not in the range
    assert not is_cell_in_merge_interior(2, 1, format_config)


def test_merge_interior_multi_row():
    format_config = {"merged_cells": ["A1:B3"]}
    assert not is_cell_in_merge_interior(1, 1, format_config)   # top-left
    assert is_cell_in_merge_interior(1, 2, format_config)        # interior col
    assert is_cell_in_merge_interior(2, 1, format_config)        # interior row
    assert is_cell_in_merge_interior(3, 2, format_config)        # interior row+col


# ---------------------------------------------------------------------------
# export — merged cells
# ---------------------------------------------------------------------------

class _FakeQuestion:
    def __init__(self, id, row_index, question_text, draft_answer=None, final_answer=None):
        self.id = id
        self.row_index = row_index
        self.question_text = question_text
        self.draft_answer = draft_answer
        self.final_answer = final_answer


class _FakeQuestionnaire:
    def __init__(self, raw_xlsx_blob, answer_col_index=None, question_count=2):
        self.id = 1
        self.raw_xlsx_blob = raw_xlsx_blob
        self.answer_col_index = answer_col_index
        self.question_count = question_count


class _FakeDB:
    """Minimal DB stub that returns pre-set questionnaire and questions."""
    def __init__(self, questionnaire, questions):
        self._questionnaire = questionnaire
        self._questions = questions

    def query(self, model):
        return self

    def filter(self, *args, **kwargs):
        return self

    def first(self):
        return self._questionnaire

    def order_by(self, *args):
        return self

    def all(self):
        return self._questions


def test_export_skips_merge_interiors():
    """Answers should not be written into non-top-left merged cells."""
    xlsx = _make_merged_xlsx()

    # row_index=0 → xlsx row 2 (header); row_index=1 → xlsx row 3; row_index=2 → row 4
    questions = [
        _FakeQuestion(1, 1, "Do you encrypt data?", draft_answer="Yes"),
        _FakeQuestion(2, 2, "Do you have MFA?", draft_answer="Yes"),
    ]
    questionnaire = _FakeQuestionnaire(raw_xlsx_blob=xlsx, answer_col_index=1, question_count=2)
    db = _FakeDB(questionnaire, questions)

    result = export_questionnaire_to_xlsx(1, db)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    ws = wb.active

    # The merged cell A1:B1 — B1 (row=1,col=2) should not have been written
    # (it's an interior merge cell; the answer rows are at rows 3 and 4)
    assert ws.cell(row=3, column=2).value == "Yes"
    assert ws.cell(row=4, column=2).value == "Yes"


def test_export_skips_formula_cells():
    """Formula cells in the answer column must not be overwritten."""
    xlsx = _make_formula_xlsx()

    questions = [
        _FakeQuestion(1, 0, "Do you encrypt data?", draft_answer="Yes"),
    ]
    questionnaire = _FakeQuestionnaire(raw_xlsx_blob=xlsx, answer_col_index=1, question_count=1)
    db = _FakeDB(questionnaire, questions)

    result = export_questionnaire_to_xlsx(1, db)
    wb = openpyxl.load_workbook(io.BytesIO(result), data_only=False)
    ws = wb.active
    # Cell B2 had "=A2"; it should not be replaced with our answer
    assert ws.cell(row=2, column=2).value == "=A2"


def test_export_appends_to_new_column_when_no_answer_col():
    """When answer_col_index is None, answers go in the column after the last used."""
    xlsx = _make_simple_xlsx(["Q1", "Q2"])

    questions = [
        _FakeQuestion(1, 0, "Q1", draft_answer="Ans1"),
        _FakeQuestion(2, 1, "Q2", draft_answer="Ans2"),
    ]
    questionnaire = _FakeQuestionnaire(raw_xlsx_blob=xlsx, answer_col_index=None, question_count=2)
    db = _FakeDB(questionnaire, questions)

    result = export_questionnaire_to_xlsx(1, db)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    ws = wb.active
    # max_column was 2 (Question, Answer header) → answers written to col 3
    assert ws.cell(row=2, column=3).value == "Ans1"
    assert ws.cell(row=3, column=3).value == "Ans2"
