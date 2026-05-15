"""Tests for the parsing service."""

import io

import openpyxl
import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.parsing import (
    parse_docx,
    parse_file,
    parse_pdf,
    parse_xlsx_preview,
    parse_xlsx_qa,
)


# ---------------------------------------------------------------------------
# Helpers — build fixture bytes in-memory using the same libraries
# ---------------------------------------------------------------------------

def _make_pdf_bytes(text: str = "") -> bytes:
    """Minimal single-page PDF, optionally with a text content stream."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    if text:
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_data = f"BT /F1 12 Tf 72 700 Td ({safe}) Tj ET".encode()
        stream_obj = DecodedStreamObject()
        stream_obj.set_data(stream_data)
        content_ref = writer._add_object(stream_obj)  # noqa: SLF001

        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)  # noqa: SLF001
        page[NameObject("/Contents")] = content_ref
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_docx_bytes(*paragraphs: str) -> bytes:
    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# parse_pdf
# ---------------------------------------------------------------------------

class TestParsePDF:
    def test_valid_blank_pdf_returns_string(self):
        result = parse_pdf(_make_pdf_bytes())
        assert isinstance(result, str)

    def test_corrupted_bytes_raises(self):
        with pytest.raises(ValueError, match="Could not read PDF"):
            parse_pdf(b"this is not a pdf")

    def test_empty_bytes_raises(self):
        with pytest.raises(ValueError):
            parse_pdf(b"")


# ---------------------------------------------------------------------------
# parse_docx
# ---------------------------------------------------------------------------

class TestParseDocx:
    def test_extracts_single_paragraph(self):
        result = parse_docx(_make_docx_bytes("Security policy text here."))
        assert "Security policy text here." in result

    def test_extracts_multiple_paragraphs(self):
        result = parse_docx(_make_docx_bytes("First paragraph.", "Second paragraph."))
        assert "First paragraph." in result
        assert "Second paragraph." in result

    def test_corrupted_bytes_raises(self):
        with pytest.raises(ValueError, match="Could not read DOCX"):
            parse_docx(b"not a valid docx")

    def test_empty_paragraphs_are_skipped(self):
        result = parse_docx(_make_docx_bytes("Real content.", "", "   "))
        assert result.strip() == "Real content."


# ---------------------------------------------------------------------------
# parse_xlsx_preview
# ---------------------------------------------------------------------------

class TestParseXLSXPreview:
    def test_returns_expected_structure(self):
        data = [
            ["Question", "Answer", "Notes"],
            ["Do you have MFA?", "Yes", "Admin accounts only"],
            ["Password rotation?", "90 days", ""],
        ]
        result = parse_xlsx_preview(_make_xlsx_bytes(data))
        assert result["columns"] == ["Question", "Answer", "Notes"]
        assert result["total_rows"] == 2
        assert len(result["sample_rows"]) == 2

    def test_sample_capped_at_five(self):
        data = [["Q", "A"]] + [[f"Q{i}", f"A{i}"] for i in range(10)]
        result = parse_xlsx_preview(_make_xlsx_bytes(data))
        assert len(result["sample_rows"]) == 5
        assert result["total_rows"] == 10

    def test_empty_sheet(self):
        result = parse_xlsx_preview(_make_xlsx_bytes([]))
        assert result == {"columns": [], "sample_rows": [], "total_rows": 0}

    def test_none_header_cell_gets_placeholder(self):
        data = [["Question", None, "Notes"], ["q1", "a1", "n1"]]
        result = parse_xlsx_preview(_make_xlsx_bytes(data))
        assert result["columns"][1] == "Column 2"

    def test_corrupted_bytes_raises(self):
        with pytest.raises(ValueError, match="Could not read XLSX"):
            parse_xlsx_preview(b"not xlsx")


# ---------------------------------------------------------------------------
# parse_xlsx_qa
# ---------------------------------------------------------------------------

class TestParseXLSXQA:
    def test_returns_correct_pairs(self):
        data = [
            ["Question", "Answer"],
            ["Do you have MFA?", "Yes, for all admins."],
            ["Password policy?", "90-day rotation enforced."],
        ]
        result = parse_xlsx_qa(_make_xlsx_bytes(data), question_col=0, answer_col=1)
        assert len(result) == 2
        assert result[0] == {"question": "Do you have MFA?", "answer": "Yes, for all admins."}

    def test_skips_empty_questions(self):
        data = [
            ["Q", "A"],
            ["Real question?", "Real answer."],
            ["", "Orphan answer"],
            [None, "Another orphan"],
        ]
        result = parse_xlsx_qa(_make_xlsx_bytes(data), question_col=0, answer_col=1)
        assert len(result) == 1

    def test_none_answer_becomes_empty_string(self):
        data = [["Q", "A"], ["Question?", None]]
        result = parse_xlsx_qa(_make_xlsx_bytes(data), question_col=0, answer_col=1)
        assert result[0]["answer"] == ""

    def test_empty_sheet_returns_empty_list(self):
        result = parse_xlsx_qa(_make_xlsx_bytes([]), question_col=0, answer_col=1)
        assert result == []

    def test_non_default_columns(self):
        data = [
            ["ID", "Category", "Question text", "Answer text"],
            ["1", "Security", "Is encryption used?", "Yes, AES-256."],
        ]
        result = parse_xlsx_qa(_make_xlsx_bytes(data), question_col=2, answer_col=3)
        assert result[0]["question"] == "Is encryption used?"
        assert result[0]["answer"] == "Yes, AES-256."


# ---------------------------------------------------------------------------
# parse_file dispatcher
# ---------------------------------------------------------------------------

class TestParseFile:
    def test_dispatches_pdf(self):
        result = parse_file("policy.pdf", _make_pdf_bytes())
        assert isinstance(result, str)

    def test_dispatches_docx(self):
        result = parse_file("policy.docx", _make_docx_bytes("Policy content."))
        assert "Policy content." in result

    def test_xlsx_raises_informative_error(self):
        with pytest.raises(ValueError, match="xlsx"):
            parse_file("data.xlsx", b"anything")

    def test_unsupported_extension_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            parse_file("data.csv", b"a,b,c")
